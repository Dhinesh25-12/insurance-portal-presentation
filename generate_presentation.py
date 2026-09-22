#!/usr/bin/env python3
"""
Insurance Portal Application - PowerPoint Presentation Generator
Generates a 10-slide professional presentation
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

def create_presentation():
    # Create presentation object
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Color scheme
    PRIMARY_COLOR = RGBColor(0, 102, 204)
    SECONDARY_COLOR = RGBColor(51, 51, 51)
    ACCENT_COLOR = RGBColor(255, 102, 0)
    WHITE = RGBColor(255, 255, 255)
    DARK_GRAY = RGBColor(51, 51, 51)
    LIGHT_GRAY = RGBColor(240, 240, 240)
    
    # Slide 1: Title Slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = PRIMARY_COLOR
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Insurance Portal Application"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(2))
    tf = subtitle_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Final Demo & Findings\nComplete Showcase of Implemented Features, Architecture & Deployment"
    p.font.size = Pt(24)
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # Slide 2: Project Overview
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    
    # Title bar
    title_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.8))
    title_shape.fill.solid()
    title_shape.fill.fore_color.rgb = PRIMARY_COLOR
    title_shape.line.color.rgb = PRIMARY_COLOR
    tf = title_shape.text_frame
    p = tf.paragraphs[0]
    p.text = "Project Overview"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # Content
    content = [
        "✓ Comprehensive Insurance Portal Application - End-to-end solution",
        "✓ Modern Technology Stack - Java backend with TypeScript frontend",
        "✓ Scalable Architecture - Microservices design for independent scaling",
        "✓ Secure & User-Friendly - Role-based access control and intuitive UI",
        "✓ Cloud-Ready Deployment - Containerized with Docker and Infrastructure as Code",
        "✓ Production Timeline - Completed in 16 days from conception to final demo"
    ]
    
    content_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.1), Inches(8.6), Inches(6))
    tf = content_box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(content):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_GRAY
        p.space_before = Pt(8)
        p.space_after = Pt(8)
    
    # Slide 3: Business Modules Implemented
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    
    title_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.8))
    title_shape.fill.solid()
    title_shape.fill.fore_color.rgb = PRIMARY_COLOR
    title_shape.line.color.rgb = PRIMARY_COLOR
    tf = title_shape.text_frame
    p = tf.paragraphs[0]
    p.text = "Business Modules Implemented"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    content = [
        "1. Policy Management - Creation, renewal, modification, and tracking",
        "2. Claims Processing - Full lifecycle from submission to settlement",
        "3. Customer Profile Management - Centralized customer information hub",
        "4. Payment & Billing - Secure payment processing and invoice management",
        "5. Document Management - Digital storage and retrieval system",
        "6. Reporting & Analytics - Business dashboards and insights",
        "7. User Authentication & Authorization - Role-based access control"
    ]
    
    content_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.1), Inches(8.6), Inches(6))
    tf = content_box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(content):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(16)
        p.font.color.rgb = DARK_GRAY
        p.space_before = Pt(6)
        p.space_after = Pt(6)
    
    # Slide 4: Application Architecture
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    
    title_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.8))
    title_shape.fill.solid()
    title_shape.fill.fore_color.rgb = PRIMARY_COLOR
    title_shape.line.color.rgb = PRIMARY_COLOR
    tf = title_shape.text_frame
    p = tf.paragraphs[0]
    p.text = "Application Architecture"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    content = [
        "Frontend Layer (UI Repository - TypeScript 56.7%)",
        "  • React-based single-page application",
        "  • HTML5 semantic markup (16.2%) and SCSS styling (10.7%)",
        "  • Responsive design optimized for all devices",
        "",
        "Backend Layer (Backend Repository - Java 82.3%)",
        "  • Spring Boot microservices framework",
        "  • RESTful APIs for seamless client-server communication",
        "  • JPA/Hibernate for data persistence",
        "",
        "Infrastructure (HCL 14.9%, Shell 2.4%, Docker 0.4%)",
        "  • Docker containerization for deployment",
        "  • Terraform/HCL for Infrastructure as Code"
    ]
    
    content_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.1), Inches(8.6), Inches(6))
    tf = content_box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(content):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(15)
        p.font.color.rgb = DARK_GRAY
        p.space_before = Pt(4)
        p.space_after = Pt(4)
    
    # Slide 5: Technology Stack
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    
    title_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.8))
    title_shape.fill.solid()
    title_shape.fill.fore_color.rgb = PRIMARY_COLOR
    title_shape.line.color.rgb = PRIMARY_COLOR
    tf = title_shape.text_frame
    p = tf.paragraphs[0]
    p.text = "Technology Stack"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    content = [
        "Frontend Technologies:",
        "  • TypeScript (56.7%) - Type-safe JavaScript programming language",
        "  • React - Component-based UI framework for dynamic interfaces",
        "  • HTML5 (16.2%) - Semantic markup for accessibility",
        "  • SCSS (10.7%) - Advanced styling with variables and mixins",
        "",
        "Backend Technologies:",
        "  • Java (82.3%) - Enterprise-grade language for scalability",
        "  • Spring Boot - Rapid microservices development framework",
        "  • JPA/Hibernate - Object-relational mapping for data access",
        "  • MySQL/PostgreSQL - Relational database management",
        "",
        "DevOps & Infrastructure:",
        "  • Docker - Containerization for consistent environments",
        "  • Terraform/HCL (14.9%) - Infrastructure as Code automation"
    ]
    
    content_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.1), Inches(8.6), Inches(6))
    tf = content_box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(content):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(14)
        p.font.color.rgb = DARK_GRAY
        p.space_before = Pt(3)
        p.space_after = Pt(3)
    
    # Slide 6: Technical Decisions & Rationale
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    
    title_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.8))
    title_shape.fill.solid()
    title_shape.fill.fore_color.rgb = PRIMARY_COLOR
    title_shape.line.color.rgb = PRIMARY_COLOR
    tf = title_shape.text_frame
    p = tf.paragraphs[0]
    p.text = "Technical Decisions & Rationale"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    content = [
        "✓ Java Backend - Enterprise support, scalability, mature ecosystem",
        "✓ TypeScript Frontend - Type safety, maintainability, reduced runtime errors",
        "✓ Microservices Architecture - Independent scaling, agile deployments",
        "✓ RESTful APIs - Industry standard, easy testing, wide framework support",
        "✓ Docker Containers - Environment consistency, rapid deployment",
        "✓ Infrastructure as Code (Terraform) - Reproducible, version-controlled infrastructure",
        "✓ Relational Database - ACID compliance, data integrity, proven reliability",
        "✓ Spring Boot Framework - Rapid development, built-in security features"
    ]
    
    content_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.1), Inches(8.6), Inches(6))
    tf = content_box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(content):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(16)
        p.font.color.rgb = DARK_GRAY
        p.space_before = Pt(6)
        p.space_after = Pt(6)
    
    # Slide 7: Deployment Approach
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    
    title_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.8))
    title_shape.fill.solid()
    title_shape.fill.fore_color.rgb = PRIMARY_COLOR
    title_shape.line.color.rgb = PRIMARY_COLOR
    tf = title_shape.text_frame
    p = tf.paragraphs[0]
    p.text = "Deployment Approach"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    content = [
        "CI/CD Pipeline:",
        "  • Automated build, test, and deployment processes",
        "  • GitHub Actions for workflow automation",
        "  • Automated testing at each stage",
        "",
        "Containerization Strategy:",
        "  • Docker for consistent application packaging",
        "  • Multi-stage builds for optimized images",
        "  • Container registry for centralized image management",
        "",
        "Infrastructure as Code:",
        "  • Terraform/HCL (14.9%) for cloud resource provisioning",
        "  • Version-controlled infrastructure configurations",
        "",
        "Environment Strategy:",
        "  • Development, Staging, and Production environments",
        "  • Environment parity through containerization"
    ]
    
    content_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.1), Inches(8.6), Inches(6))
    tf = content_box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(content):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(14)
        p.font.color.rgb = DARK_GRAY
        p.space_before = Pt(3)
        p.space_after = Pt(3)
    
    # Slide 8: Challenges Encountered & Solutions
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    
    title_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.8))
    title_shape.fill.solid()
    title_shape.fill.fore_color.rgb = PRIMARY_COLOR
    title_shape.line.color.rgb = PRIMARY_COLOR
    tf = title_shape.text_frame
    p = tf.paragraphs[0]
    p.text = "Challenges Encountered & Solutions"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    content = [
        "Challenge 1: Microservices Communication",
        "  ✓ Solution: Implemented API gateway pattern with proper service contracts",
        "",
        "Challenge 2: Database Performance",
        "  ✓ Solution: Strategic indexing, query optimization, caching layer",
        "",
        "Challenge 3: Frontend-Backend Integration",
        "  ✓ Solution: Comprehensive API documentation and API mocking services",
        "",
        "Challenge 4: Security & Compliance",
        "  ✓ Solution: JWT authentication, encryption, audit logging, RBAC",
        "",
        "Challenge 5: Deployment Consistency",
        "  ✓ Solution: Docker containerization and Terraform automation"
    ]
    
    content_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.1), Inches(8.6), Inches(6))
    tf = content_box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(content):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(15)
        p.font.color.rgb = DARK_GRAY
        p.space_before = Pt(4)
        p.space_after = Pt(4)
    
    # Slide 9: Lessons Learned
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    
    title_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.8))
    title_shape.fill.solid()
    title_shape.fill.fore_color.rgb = PRIMARY_COLOR
    title_shape.line.color.rgb = PRIMARY_COLOR
    tf = title_shape.text_frame
    p = tf.paragraphs[0]
    p.text = "Lessons Learned"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    content = [
        "✓ Early API Contract Definition - Reduces integration issues significantly",
        "✓ Containerization Benefits - Simplifies deployment and environment consistency",
        "✓ Infrastructure as Code - Ensures repeatability, scalability, and disaster recovery",
        "✓ Type Safety Matters - TypeScript catches errors early in development",
        "✓ Monitoring is Critical - Microservices require robust logging and observability",
        "✓ Communication is Key - Team coordination essential in distributed systems",
        "✓ Testing Coverage - Comprehensive testing saves debugging time",
        "✓ Documentation Importance - Well-documented APIs accelerate development",
        "✓ Agile Deployment - Frequent small releases reduce deployment risk"
    ]
    
    content_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.1), Inches(8.6), Inches(6))
    tf = content_box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(content):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(15)
        p.font.color.rgb = DARK_GRAY
        p.space_before = Pt(5)
        p.space_after = Pt(5)
    
    # Slide 10: Future Enhancements & Roadmap
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    
    title_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.8))
    title_shape.fill.solid()
    title_shape.fill.fore_color.rgb = PRIMARY_COLOR
    title_shape.line.color.rgb = PRIMARY_COLOR
    tf = title_shape.text_frame
    p = tf.paragraphs[0]
    p.text = "Future Enhancements & Roadmap"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    content = [
        "Scalability Enhancements:",
        "  • Horizontal scaling for handling high-traffic scenarios",
        "  • Database sharding for improved performance",
        "  • Caching layers (Redis) for frequently accessed data",
        "",
        "New Features:",
        "  • AI-powered claims assessment and fraud detection",
        "  • Mobile applications for iOS and Android",
        "  • Advanced analytics and predictive modeling",
        "  • Real-time notifications and mobile push alerts",
        "",
        "Performance & Security:",
        "  • GraphQL API for optimized data fetching",
        "  • Zero-trust security architecture",
        "  • Advanced threat detection and prevention",
        "  • Compliance with additional regulatory standards"
    ]
    
    content_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.1), Inches(8.6), Inches(6))
    tf = content_box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(content):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(14)
        p.font.color.rgb = DARK_GRAY
        p.space_before = Pt(3)
        p.space_after = Pt(3)
    
    return prs

if __name__ == "__main__":
    prs = create_presentation()
    output_file = "Insurance_Portal_Application_Demo.pptx"
    prs.save(output_file)
    print(f"✅ PowerPoint presentation created successfully!")
    print(f"📁 File: {output_file}")
    print(f"📊 Slides: 10 comprehensive slides")
    print(f"✨ Ready for download and presentation!")
